"""Fill the AI Builder Cup template with the Second Voice pitch.

Keeps every template slide's header art and heading text box, adds content below the
heading, deletes the instruction slide, the optional wireframes slide and the blank slide.
Run: uvx --from python-pptx python -I build_deck.py template.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC, OUT = sys.argv[1], sys.argv[2]
ASSETS = "deck-assets"
FONT = "Arial"
INK = RGBColor(0x16, 0x20, 0x2B)
MUTED = RGBColor(0x5B, 0x67, 0x74)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_SOFT = RGBColor(0xFD, 0xF1, 0xE2)
PANEL = RGBColor(0xF3, 0xF4, 0xF6)
LINE = RGBColor(0xD9, 0xDE, 0xE5)
GREEN = RGBColor(0x15, 0x80, 0x3D)
RED = RGBColor(0xB9, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LINKS = {
    "github": "https://github.com/R-I-T-I-K/second-voice",
    "video": "Demo video link: to be added after upload",
    "product": "https://fabled-ray-512016-k5.web.app/",
}

prs = Presentation(SRC)
TOP = Inches(1.45)          # content starts under the template heading
BOTTOM = Inches(5.35)       # above the template's footer line
LEFT = Inches(0.35)
WIDTH = Inches(9.3)


# ---------- helpers ----------

def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, font=FONT, margin=0.05):
    """paras: list of str or (str, {opts}) or list-of-runs [(str, opts), ...]."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(margin))
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        para.space_after = Pt(space_after)
        opts = {}
        runs = p
        if isinstance(p, tuple):
            runs, opts = p
        if isinstance(runs, str):
            runs = [(runs, {})]
        for rt, ro in runs:
            r = para.add_run()
            r.text = rt
            f = r.font
            f.name = ro.get("font", opts.get("font", font))
            f.size = Pt(ro.get("size", opts.get("size", size)))
            f.bold = ro.get("bold", opts.get("bold", bold))
            f.italic = ro.get("italic", opts.get("italic", False))
            f.color.rgb = ro.get("color", opts.get("color", color))
        if opts.get("bullet"):
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.18)))
            pPr.set("indent", str(-Inches(0.18)))
            from pptx.oxml.ns import qn
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    if radius:
        shp.adjustments[0] = 0.06
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def card(slide, x, y, w, h, title, body, fill=PANEL, title_color=INK, body_size=10.5, title_size=12.5, tag=None):
    box(slide, x, y, w, h, fill=fill)
    cy = y + Inches(0.12)
    if tag:
        text(slide, x + Inches(0.15), cy, w - Inches(0.3), Inches(0.22), [(tag, {"size": 8, "color": AMBER, "bold": True})],
             margin=0.02)
        cy += Inches(0.22)
    text(slide, x + Inches(0.15), cy, w - Inches(0.3), Inches(0.35), [(title, {"size": title_size, "bold": True, "color": title_color})], margin=0.02)
    text(slide, x + Inches(0.15), cy + Inches(0.36), w - Inches(0.3), h - (cy - y) - Inches(0.42),
         [(b, {"size": body_size, "color": INK}) if isinstance(b, str) else b for b in body], space_after=3, margin=0.02)


def arrow(slide, x, y, w=Inches(0.28), h=Inches(0.28), color=MUTED):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()
    a.shadow.inherit = False
    return a


def picture_fit(slide, path, x, y, max_w, max_h, border=True):
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(path, x, y, w, h)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic, w, h


def set_heading(slide, new_text):
    """Replace the template heading text while keeping its run formatting."""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = new_text
                for r in p.runs[1:]:
                    r.text = ""
            # drop extra template paragraphs (the prompt bullets)
            for extra in sh.text_frame.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            return sh
    return None


slides = list(prs.slides)
S = {i + 1: s for i, s in enumerate(slides)}

# ---------- slide 1: team details ----------
s = S[1]
for sh in s.shapes:
    if sh.has_text_frame and "Team Details" in sh.text_frame.text:
        values = {"Team name:": " Second Voice", "Team leader name:": " Ritik Singh",
                  "Problem Statement:": " BFSI. AI to detect and stop fraud that bypasses today's controls: authorised push payment "
                                        "(coercion) scams on UPI, such as digital arrest."}
        for p in sh.text_frame.paragraphs:
            for key, val in values.items():
                if p.text.strip().startswith(key) and p.runs:
                    p.runs[-1].text = p.runs[-1].text.rstrip() + val

# ---------- slide 3: brief ----------
s = S[3]
set_heading(s, "Brief about the idea")
text(s, LEFT, TOP, Inches(5.3), Inches(0.6),
     [[("Second Voice", {"size": 22, "bold": True}), ("  is a coercion-aware payment interruptor for UPI.", {"size": 14, "color": MUTED})]])
text(s, LEFT, TOP + Inches(0.75), Inches(5.3), Inches(3.2), [
    ("Every bank fraud system asks one question: is this really the account holder? In digital arrest, parcel and "
     "investment scams the answer is yes. The victim is on her own phone, typing her own PIN, because a caller posing as "
     "a CBI officer told her to. Device, location and credential checks see nothing.", {"size": 11.5}),
    ("Second Voice sits in the payer bank's authorisation hook, before the debit. It scores whether the account holder is "
     "being manipulated right now, from eight signals around the moment of payment, and answers with a Gemini-driven "
     "conversation built from the suspected scam script, not a warning banner.", {"size": 11.5}),
    ("The receiving bank runs the same service on its own accounts and shares one word back during address validation: "
     "a receiver risk band. If the customer overrides anyway, the receiving bank holds the credit before it can be cashed out.",
     {"size": 11.5}),
], space_after=7)
# right: the gap
box(s, Inches(5.95), TOP, Inches(3.7), Inches(3.75), fill=PANEL)
text(s, Inches(6.1), TOP + Inches(0.12), Inches(3.4), Inches(0.3), [("THE GAP TODAY", {"size": 8.5, "bold": True, "color": AMBER})], margin=0.02)
text(s, Inches(6.1), TOP + Inches(0.42), Inches(3.4), Inches(1.5), [
    ("Account takeover", {"size": 12, "bold": True}),
    ("Stolen credentials, new device, odd location. Detected by every bank today.", {"size": 10.5, "color": MUTED}),
], space_after=3, margin=0.02)
text(s, Inches(6.1), TOP + Inches(1.5), Inches(3.4), Inches(1.6), [
    ("Authorised push payment fraud", {"size": 12, "bold": True, "color": RED}),
    ("Real customer, real PIN, under coercion. Looks legitimate to every rule. The only intervention in production is a "
     "static banner, dismissed in two seconds.", {"size": 10.5, "color": MUTED}),
], space_after=3, margin=0.02)
text(s, Inches(6.1), TOP + Inches(2.85), Inches(3.4), Inches(0.8), [
    ("Rising fast in India (digital arrest, task and investment scams). The UK now makes sending and receiving banks "
     "reimburse victims 50:50.", {"size": 10, "italic": True, "color": INK}),
], margin=0.02)

# ---------- slide 4: opportunities ----------
s = S[4]
set_heading(s, "Opportunities")
cols = [
    ("HOW IT IS DIFFERENT", "Models coercion, not takeover",
     ["Existing fraud engines model account takeover. Second Voice models coercion: payee novelty, amount against the "
      "customer's own history, an active call, a remote-access app, fan-in from our other customers, and a band from the receiving bank.",
      "The response is a conversation shaped by the suspected scam script, with questions the scammer did not coach the victim for."]),
    ("HOW IT SOLVES THE PROBLEM", "Two banks, two chances",
     ["Payer side: pause before the debit, ask, offer a trusted-contact call and a cooling-off, then release, hold or cancel.",
      "Receiver side: the beneficiary bank scores its own accounts (age, KYC depth, first-time inflow fan-in, sweep-out speed) "
      "and holds credits to High-band accounts, so money stays inside the banking system even after an override."]),
    ("USP", "A network, not a rule",
     ["One field added to the existing UPI address-validation lookup carries the receiver band. Every bank that joins sharpens "
      "every other bank's score.",
      "Graduated tiers mean low-risk payments see nothing. Fraud teams get the full reasoning, transcript and a Gemini summary, "
      "plus a payee blocklist and I4C report in one click."]),
]
cw = Inches(3.0)
for i, (tag, title, body) in enumerate(cols):
    x = LEFT + i * (cw + Inches(0.15))
    card(s, x, TOP, cw, Inches(3.75), title, body, tag=tag, body_size=11)

# ---------- slide 5: features ----------
s = S[5]
set_heading(s, "List of features offered by the solution")
feats = [
    ("Coercion score", "Weighted model over eight interpretable signals, each shown with its contribution. Tiers: low, medium, high, very high."),
    ("Script-aware intervention", "Gemini names the scam script (digital arrest, parcel, KYC expiry, investment group, task fraud) and runs an adaptive conversation."),
    ("Graduated friction", "Low: nothing. Medium: one question. High: hold, trusted-contact call, cooling-off. Very high: block and route to a specialist."),
    ("Receiver risk band", "The receiving bank scores its own accounts and answers address validation with Low, Elevated or High. No account data crosses the wall."),
    ("Credit hold at the receiver", "An override still lands on a 24-hour hold at the beneficiary bank when the account is High band."),
    ("Fraud operations", "Live case queue, signals, transcript with the agent's reasoning, Gemini case summary, one-click payee blocklist and I4C report."),
]
cw, ch = Inches(3.0), Inches(1.75)
for i, (t, b) in enumerate(feats):
    x = LEFT + (i % 3) * (cw + Inches(0.15))
    y = TOP + (i // 3) * (ch + Inches(0.15))
    card(s, x, y, cw, ch, t, [b], body_size=10)

# ---------- slide 6: process flow ----------
s = S[6]
set_heading(s, "Process flow")
row1_y = TOP + Inches(0.1)
bw, bh, gap = Inches(1.55), Inches(1.0), Inches(0.3)
steps = [
    ("1  Tap Pay", "Bank A app. SDK attaches device context."),
    ("2  Address check", "Bank A asks Bank B for the name. Bank B answers with name + risk band."),
    ("3  Coercion score", "History, payee reputation, device, band, registry. Eight signals, one score."),
    ("4  Tier", "Low: allow. Medium/High: hold and talk. Very high: block."),
    ("5  Gemini talks", "Script-aware questions. Re-scored on every turn."),
]
x = LEFT
for i, (t, b) in enumerate(steps):
    box(s, x, row1_y, bw, bh, fill=PANEL if i != 4 else AMBER_SOFT)
    text(s, x + Inches(0.08), row1_y + Inches(0.06), bw - Inches(0.16), Inches(0.3), [(t, {"size": 10, "bold": True})], margin=0.02)
    text(s, x + Inches(0.08), row1_y + Inches(0.36), bw - Inches(0.16), bh - Inches(0.4), [(b, {"size": 8.5, "color": MUTED})], margin=0.02)
    if i < len(steps) - 1:
        arrow(s, x + bw + Inches(0.02), row1_y + bh / 2 - Inches(0.12), w=Inches(0.26), h=Inches(0.24))
    x += bw + gap
# outcomes row
row2_y = row1_y + bh + Inches(0.55)
text(s, LEFT, row1_y + bh + Inches(0.12), WIDTH, Inches(0.3), [("OUTCOMES", {"size": 8.5, "bold": True, "color": AMBER})], margin=0.02)
outs = [
    ("Release", "Customer knows the payee. Payment completes.", GREEN),
    ("Trusted-contact hold", "Daughter is called. 24-hour pause. Nothing moves.", AMBER),
    ("Cancel", "Customer cancels. Case closed, payee blocklisted if confirmed.", INK),
    ("Override", "10-second cooling-off, then debit. Bank B holds the credit if High band.", RED),
    ("Fraud ops", "Transcript, summary, outcome label, I4C report. Outcomes retrain the score.", MUTED),
]
ow = Inches(1.75)
x = LEFT
for t, b, c in outs:
    box(s, x, row2_y, ow, Inches(1.25), fill=WHITE, line=LINE)
    text(s, x + Inches(0.1), row2_y + Inches(0.08), ow - Inches(0.2), Inches(0.3), [(t, {"size": 10.5, "bold": True, "color": c})], margin=0.02)
    text(s, x + Inches(0.1), row2_y + Inches(0.4), ow - Inches(0.2), Inches(0.8), [(b, {"size": 8.5, "color": MUTED})], margin=0.02)
    x += ow + Inches(0.13)

# ---------- slide 8: architecture ----------
s = S[8]
set_heading(s, "Architecture")
pic, pw, ph = picture_fit(s, f"{ASSETS}/architecture.png", LEFT, TOP - Inches(0.1), Inches(5.5), Inches(4.05), border=False)
tx = LEFT + pw + Inches(0.25)
tw = Inches(9.65) - tx
text(s, tx, TOP - Inches(0.05), tw, Inches(4.0), [
    ("One codebase, two banks", {"size": 13, "bold": True}),
    ("The same Cloud Run service is deployed as bank-a (payer hook) and bank-b (receiver band). Firestore is the only state; "
     "live listeners drive the phone and both consoles.", {"size": 10, "color": MUTED}),
    ("Flow", {"size": 13, "bold": True}),
    ("1  App sends pay request + device context", {"size": 9.5, "bullet": True}),
    ("2  Payment API validates the address with Bank B, gets the band", {"size": 9.5, "bullet": True}),
    ("3  Scorer joins history, fan-in, band, registry; Gemini names the script", {"size": 9.5, "bullet": True}),
    ("4  Medium and above: Gemini intervention agent, turns in Firestore", {"size": 9.5, "bullet": True}),
    ("5  Decision persists; ops console updates live", {"size": 9.5, "bullet": True}),
    ("6  Override: credit to Bank B, held if High band", {"size": 9.5, "bullet": True}),
    ("7  Outcomes labelled by ops feed back into the score", {"size": 9.5, "bullet": True}),
    ("Dashed boxes are simulated in the prototype: device signals (toggles), registry lookups, the NPCI leg, the trusted-contact call.",
     {"size": 9, "italic": True, "color": MUTED}),
], space_after=4)

# ---------- slide 9: technologies ----------
s = S[9]
set_heading(s, "Technologies used in the solution")
techs = [
    ("Cloud Run", "Both bank services from one container. Payment API, scorer, Gemini agent, address validation, credit hold. One minimum instance each, asia-south1.", True),
    ("Firestore", "Customers, 90-day history, receiver accounts, cases, holds, conversation turns. Composite indexes for the fan-in queries. Live listeners to every screen.", True),
    ("Firebase Hosting + Auth", "Mock bank app and two ops consoles as static sites. Anonymous auth gates Firestore reads; all writes come from Cloud Run.", True),
    ("Gemini 2.5 Flash on Vertex AI", "Names the scam script from the signal pattern, runs the intervention turns with a per-script system prompt, writes the case summary. JSON-mode responses.", True),
    ("FastAPI, Python 3.12, uv", "Service code, weighted scoring model, receiver mule score, Firestore Admin SDK, httpx between banks.", False),
    ("Firebase JS SDK 10", "Phone app and consoles in plain HTML/JS; onSnapshot listeners for cases, turns, accounts, holds, events.", False),
    ("Cloud Build, Artifact Registry", "Source-to-container deploys through gcloud run deploy --source.", False),
    ("Planned", "BigQuery analytics copy of history and outcomes, Pub/Sub event stream, Vertex AI Search over the scam-script corpus, trained model from labelled cases.", False),
]
cw, ch = Inches(4.55), Inches(0.9)
for i, (t, b, core) in enumerate(techs):
    x = LEFT + (i % 2) * (cw + Inches(0.2))
    y = TOP + (i // 2) * (ch + Inches(0.1))
    box(s, x, y, cw, ch, fill=AMBER_SOFT if core else PANEL)
    text(s, x + Inches(0.12), y + Inches(0.06), cw - Inches(0.24), Inches(0.3), [(t, {"size": 11, "bold": True})], margin=0.02)
    text(s, x + Inches(0.12), y + Inches(0.34), cw - Inches(0.24), ch - Inches(0.38), [(b, {"size": 8.8, "color": MUTED})], margin=0.02)

# ---------- slide 10: cost ----------
s = S[10]
set_heading(s, "Estimated implementation cost")
text(s, LEFT, TOP, WIDTH, Inches(0.4), [("Prototype as deployed (per month, indicative, asia-south1 list prices)", {"size": 12, "bold": True})])
rows = [
    ("Cloud Run, 2 services, 1 min instance each, 512 MB", "~ $14 to $20"),
    ("Firestore, demo-scale reads and writes", "within free tier"),
    ("Gemini 2.5 Flash, ~3 calls per flagged payment", "~ $0.002 per intervention"),
    ("Firebase Hosting and Auth", "free tier"),
    ("Total", "~ $15 to $25 per month"),
]
ty = TOP + Inches(0.45)
for i, (a, b) in enumerate(rows):
    last = i == len(rows) - 1
    box(s, LEFT, ty, Inches(5.6), Inches(0.36), fill=AMBER_SOFT if last else (PANEL if i % 2 == 0 else WHITE), radius=False)
    text(s, LEFT + Inches(0.1), ty + Inches(0.03), Inches(3.3), Inches(0.3), [(a, {"size": 10, "bold": last})], margin=0.02)
    text(s, LEFT + Inches(3.4), ty + Inches(0.03), Inches(2.1), Inches(0.3), [(b, {"size": 10, "bold": last})], align=PP_ALIGN.RIGHT, margin=0.02)
    ty += Inches(0.36)
text(s, Inches(6.2), TOP, Inches(3.45), Inches(3.8), [
    ("At bank scale", {"size": 12, "bold": True}),
    ("Scoring is a few Firestore or BigQuery lookups per payment: compute cost is negligible against existing fraud-engine spend.",
     {"size": 10, "color": MUTED}),
    ("Gemini is called only for the small share of payments that reach medium tier or above. At 1 to 2 percent of volume and three "
     "calls per case, a bank doing 10 million UPI payments a month spends on the order of a few hundred dollars on the model.",
     {"size": 10, "color": MUTED}),
    ("The real cost is integration: the SDK in the bank's app, the hook in the payment switch, and the band field in address validation.",
     {"size": 10, "color": MUTED}),
    ("Set against reimbursement liability for authorised push payment fraud, which the UK now splits between sending and receiving banks.",
     {"size": 10, "italic": True}),
], space_after=6)

# ---------- slide 11: snapshots ----------
s = S[11]
set_heading(s, "Snapshots of the prototype")
from PIL import Image
def crop(src, dst, box_):
    Image.open(src).crop(box_).save(dst)
    return dst
app_img = crop(f"{ASSETS}/shot-app.png", f"{ASSETS}/crop-app.png", (280, 60, 2560, 1800))
ca_img = crop(f"{ASSETS}/shot-console-a.png", f"{ASSETS}/crop-console-a.png", (60, 60, 2820, 1010))
cb_img = crop(f"{ASSETS}/shot-console-b.png", f"{ASSETS}/crop-console-b.png", (60, 60, 2820, 1010))
lab = {"size": 9, "bold": True, "color": MUTED}
# left: phone conversation
text(s, LEFT, TOP - Inches(0.05), Inches(4.3), Inches(0.22), [("PHONE APP: PAYMENT PAUSED, GEMINI CONVERSATION, SIGNALS", lab)], margin=0.02)
_, w1, h1 = picture_fit(s, app_img, LEFT, TOP + Inches(0.18), Inches(4.3), Inches(3.3))
text(s, LEFT, TOP + Inches(0.18) + h1 + Inches(0.06), Inches(4.3), Inches(0.4),
     [("Live at fabled-ray-512016-k5.web.app. No login. Pick Digital arrest, tap Pay.", {"size": 9, "color": INK})], margin=0.02)
# right: two consoles
rx, rw = Inches(4.85), Inches(4.8)
text(s, rx, TOP - Inches(0.05), rw, Inches(0.22), [("BANK A FRAUD CONSOLE: CASE, SIGNALS, TRANSCRIPT", lab)], margin=0.02)
_, w2, h2 = picture_fit(s, ca_img, rx, TOP + Inches(0.18), rw, Inches(1.7))
y2 = TOP + Inches(0.18) + h2 + Inches(0.14)
text(s, rx, y2, rw, Inches(0.22), [("BANK B RECEIVER CONSOLE: MULE BANDS, HELD CREDITS", lab)], margin=0.02)
picture_fit(s, cb_img, rx, y2 + Inches(0.23), rw, Inches(1.7))

# ---------- slide 12: performance ----------
s = S[12]
set_heading(s, "Prototype performance report")
text(s, LEFT, TOP, Inches(4.6), Inches(0.35), [("Latency on the live deployment, 19 Oct 2026", {"size": 11.5, "bold": True})])
perf = [
    ("Service health check", "0.2 s"),
    ("Address validation, Bank A to Bank B, with band", "0.23 s"),
    ("Low-risk payment: score and allow (no model call)", "0.5 to 0.6 s"),
    ("Flagged payment: score, Gemini script classification, hold", "7.5 to 10 s"),
    ("Gemini conversation turn", "3 to 5 s"),
    ("Receiver credit hold at Bank B", "under 1 s"),
]
ty = TOP + Inches(0.4)
for i, (a, b) in enumerate(perf):
    box(s, LEFT, ty, Inches(4.6), Inches(0.36), fill=PANEL if i % 2 == 0 else WHITE, radius=False)
    text(s, LEFT + Inches(0.1), ty + Inches(0.03), Inches(3.4), Inches(0.3), [(a, {"size": 9.5})], margin=0.02)
    text(s, LEFT + Inches(3.4), ty + Inches(0.03), Inches(1.1), Inches(0.3), [(b, {"size": 9.5, "bold": True})], align=PP_ALIGN.RIGHT, margin=0.02)
    ty += Inches(0.36)
text(s, LEFT, ty + Inches(0.08), Inches(4.6), Inches(0.7), [
    ("Scoring itself is sub-second. The 7 to 10 s on flagged payments is the synchronous Gemini classification; the next revision "
     "returns the hold immediately on the rule-based script guess and lets Gemini refine it in the background.",
     {"size": 8.8, "italic": True, "color": MUTED})], margin=0.02)
# right: scoring outcomes
rx = Inches(5.2)
text(s, rx, TOP, Inches(4.45), Inches(0.35), [("Scoring outcomes on the seeded scenarios", {"size": 11.5, "bold": True})])
outs = [
    ("Sunita, digital arrest, ₹2,40,000 to a new payee on a 23-min call", "0.909", "high", AMBER),
    ("Sunita, ₹2,400 to her regular milk vendor", "0.178", "low", GREEN),
    ("New payee, 3x her usual amount, no call", "0.386", "medium", AMBER),
    ("Bank B: Rakesh Kumar, 9 days old, 12 first-time inflows, 18-min sweep", "0.917", "High", RED),
    ("Bank B: Pune Dairy, 6 years old, full KYC", "0.054", "Low", GREEN),
    ("Bank B: Meera Traders, 20 days old, 5 first-time inflows", "0.442", "Elevated", AMBER),
]
ty = TOP + Inches(0.4)
for i, (a, sc, tier, c) in enumerate(outs):
    box(s, rx, ty, Inches(4.45), Inches(0.42), fill=PANEL if i % 2 == 0 else WHITE, radius=False)
    text(s, rx + Inches(0.1), ty + Inches(0.03), Inches(3.0), Inches(0.38), [(a, {"size": 8.8})], margin=0.02)
    text(s, rx + Inches(3.1), ty + Inches(0.06), Inches(0.6), Inches(0.3), [(sc, {"size": 10, "bold": True})], align=PP_ALIGN.RIGHT, margin=0.02)
    text(s, rx + Inches(3.75), ty + Inches(0.06), Inches(0.65), Inches(0.3), [(tier, {"size": 9, "bold": True, "color": c})], align=PP_ALIGN.RIGHT, margin=0.02)
    ty += Inches(0.42)
text(s, rx, ty + Inches(0.08), Inches(4.45), Inches(0.6), [
    ("Gemini named the script correctly in every live run (digital arrest, confidence 0.95) and produced valid JSON turns and summaries; "
     "stub fallbacks never triggered.", {"size": 8.8, "italic": True, "color": MUTED})], margin=0.02)

# ---------- slide 13: future ----------
s = S[13]
set_heading(s, "Future development")
fut = [
    ("Real device signals", "Android SDK reading call state, recently installed remote-access apps and payee age, replacing the demo toggles."),
    ("Learned model", "Train the coercion score on labelled case outcomes from the ops console; BigQuery as the analytics store, Pub/Sub for the event stream."),
    ("Shared registries", "Live lookups to the I4C Suspect Registry and the DoT Financial Fraud Risk Indicator; file I4C reports automatically."),
    ("Receiver band on the rails", "Propose the risk-band field for the UPI address-validation response so any bank pair can exchange it."),
    ("Languages and voice", "Hindi and Marathi turns; a spoken intervention for customers who cannot read under pressure."),
    ("Script corpus", "Vertex AI Search over fresh scam scripts and RBI advisories, so the agent learns new patterns without a redeploy."),
]
cw, ch = Inches(3.0), Inches(1.75)
for i, (t, b) in enumerate(fut):
    x = LEFT + (i % 3) * (cw + Inches(0.15))
    y = TOP + (i // 3) * (ch + Inches(0.15))
    card(s, x, y, cw, ch, t, [b], body_size=10)

# ---------- slide 14: links ----------
s = S[14]
set_heading(s, "Links")
links = [
    ("GitHub public repository", LINKS["github"]),
    ("Demo video (3 to 4 minutes)", LINKS["video"]),
    ("Final product link", LINKS["product"]),
]
ty = TOP + Inches(0.2)
for t, u in links:
    box(s, LEFT, ty, WIDTH, Inches(0.85), fill=PANEL)
    text(s, LEFT + Inches(0.2), ty + Inches(0.1), WIDTH - Inches(0.4), Inches(0.3), [(t, {"size": 12, "bold": True})], margin=0.02)
    text(s, LEFT + Inches(0.2), ty + Inches(0.42), WIDTH - Inches(0.4), Inches(0.35), [(u, {"size": 12, "color": AMBER})], margin=0.02)
    ty += Inches(1.0)
text(s, LEFT, ty + Inches(0.1), WIDTH, Inches(0.5), [
    ("The product link needs no login. Open the Digital arrest scenario and tap Pay; the two consoles update live in other tabs.",
     {"size": 10, "color": MUTED})])

# ---------- delete slides 2 (instructions), 7 (optional wireframes), 15 (blank) ----------
from pptx.oxml.ns import qn
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
for idx in sorted([2, 7, 15], reverse=True):
    el = ids[idx - 1]
    rId = el.get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(el)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
